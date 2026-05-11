import os
from unittest.mock import MagicMock, patch

import openpyxl

from app.services.send_demand_email_service import ALL_SKUS, generate_demand_excel, send_bulk_emails, send_email


def test_generate_demand_excel_creates_expected_file(tmp_path):
    with patch.dict(os.environ, {"ATTACHMENTS_DIR": str(tmp_path)}):
        path = generate_demand_excel(["MALKIST CHEESE  48 PCS X 72"])
    assert path.endswith(".xlsx")
    assert os.path.exists(path)
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    headers = [ws.cell(1, c).value for c in range(1, 5)]
    assert headers == ["sku_id", "sku_name", "sku_description", "Quantity"]
    assert ws.max_row == len(ALL_SKUS) + 1
    assert ws.cell(2, 4).value is None


def test_send_email_success_and_missing_attachment(mock_smtp_server):
    settings = MagicMock(EMAIL_ADDRESS="sender@example.com", EMAIL_APP_PASSWORD="pw", SMTP_SERVER="smtp.example.com", SMTP_PORT=465)
    with patch("app.services.send_demand_email_service.settings", settings):
        result = send_email("to@example.com", "Subject", "Body", attachment_path="/missing/file.xlsx")
    assert result["status"] == "success"
    mock_smtp_server[1].login.assert_called_once()
    mock_smtp_server[1].send_message.assert_called_once()


def test_send_email_with_real_attachment(tmp_path, mock_smtp_server):
    attachment = tmp_path / "file.xlsx"
    attachment.write_bytes(b"demo")
    settings = MagicMock(EMAIL_ADDRESS="sender@example.com", EMAIL_APP_PASSWORD="pw", SMTP_SERVER="smtp.example.com", SMTP_PORT=465)
    with patch("app.services.send_demand_email_service.settings", settings):
        result = send_email("to@example.com", "Subject", "Body", attachment_path=str(attachment))
    assert result["status"] == "success"


def test_send_email_returns_error_on_smtp_failure():
    settings = MagicMock(EMAIL_ADDRESS="sender@example.com", EMAIL_APP_PASSWORD="pw", SMTP_SERVER="smtp.example.com", SMTP_PORT=465)
    with patch("app.services.send_demand_email_service.settings", settings), patch("smtplib.SMTP_SSL", side_effect=Exception("Connection refused")):
        result = send_email("to@example.com", "Subject", "Body")
    assert result["status"] == "error"
    assert "Connection refused" in result["message"]


def test_send_bulk_emails_success_and_failure():
    distributors = [{"id": "D01", "email": "d01@example.com"}, {"id": "D02", "email": "d02@example.com"}]
    with patch("app.services.send_demand_email_service.SKURecommendationService") as svc, \
         patch("app.services.send_demand_email_service.build_demand_email") as build, \
         patch("app.services.send_demand_email_service.generate_demand_excel", return_value="/tmp/DEMAND.xlsx"), \
         patch("app.services.send_demand_email_service.send_email", return_value={"status": "success"}):
        svc.return_value.execute.return_value = {"recommendations": ["MALKIST"]}
        build.return_value = {"subject": "Sub", "body": "Body", "to_email": "d01@example.com"}
        result = send_bulk_emails(distributors)
    assert result["status"] == "completed"
    assert len(result["results"]) == 2


def test_send_bulk_emails_handles_per_distributor_exception():
    with patch("app.services.send_demand_email_service.SKURecommendationService") as svc:
        svc.return_value.execute.side_effect = Exception("DB down")
        result = send_bulk_emails([{"id": "D01", "email": "d01@example.com"}])
    assert result["results"][0]["status"] == "error"
