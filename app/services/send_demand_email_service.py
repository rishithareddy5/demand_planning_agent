import smtplib
import ssl
import os
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart
from email.mime.base import MIMEBase
from email import encoders

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

from app.core.config import settings
from app.services.build_demand_email_service import build_demand_email
from app.services.sku_recommendation_service import SKURecommendationService


# ── All 40 SKUs (master list) ──────────────────────────────────────────────────
ALL_SKUS = [
    ("SKU01",  "MALKIST CHEESE  48 PCS X 72",        "MALKIST CHEESE  48 PCS X 72 GM- GT"),
    ("SKU02",  "MALKIST CHEESE  G",                   "MALKIST CHEESE  GB 16 X 30 X 18G"),
    ("SKU03",  "MALKIST CHEESE  FAMILY 10",           "MALKIST CHEESE  FAMILY 10 X 6 X 144GM-GT"),
    ("SKU04",  "BENG BENG WAFER 22GM",                "BENG BENG WAFER 12 X 25 X 22GM"),
    ("SKU05",  "BENG BENG WAFER GB",                  "BENG BENG WAFER GB 12 X 25 X 25G @ MRP12"),
    ("SKU06",  "KOPIKO CAPPU EXTRA",                  "KOPIKO CAPPU 12X230X3.5GR @ 15 PCS EXTRA"),
    ("SKU07",  "KOPIKO CAPPU + MALKIST",              "KOPIKO CAPPU 4X650X3.5G+MALKIST 20PCS GB"),
    ("SKU08",  "KOPIKO CAPPUCCINO",                   "KOPIKO CAPPUCCINO 28 PCH X 120 X 3.5G"),
    ("SKU09",  "MALKIST SUGAR GB",                    "MALKIST SUGAR GB 16 X 30 X 18G"),
    ("SKU010", "MALKIST SUGAR CRACKERS",              "MALKIST SUGAR CRACKERS 10X 6 X 150G-RS50"),
    ("SKU011", "MALKIST CHEESE MINI PACK",            "MALKIST CHEESE MINI 24 PCS X 36GM - IMPULSE PACK"),
    ("SKU012", "MALKIST CHEESE JUMBO PACK",           "MALKIST CHEESE JUMBO 6 PCS X 250GM - VALUE PACK"),
    ("SKU013", "MALKIST CHEESE DISPLAY BOX",          "MALKIST CHEESE DISPLAY 10 X 8 X 72GM - COUNTER DISPLAY"),
    ("SKU014", "MALKIST CHEESE GIFT BOX",             "MALKIST CHEESE GIFT 3 X 200GM - FESTIVE GIFTING PACK"),
    ("SKU015", "MALKIST CHEESE + CHOCOLATE",          "MALKIST CHEESE & CHOCO COMBO 16 X 30 X 20G - COMBO"),
    ("SKU016", "MALKIST CHEESE CRACKER THIN",         "MALKIST CHEESE THIN CRACKER 24 X 20 X 25GM - THIN FORMAT"),
    ("SKU017", "MALKIST CHEESE MULTIPACK",            "MALKIST CHEESE MULTIPACK 5 X 10 X 72GM - SCHOOL PACK"),
    ("SKU018", "MALKIST CHEESE BULK SHIPPER",         "MALKIST CHEESE SHIPPER CASE 100 PCS X 72GM - BULK TRADE"),
    ("SKU019", "BENG BENG WAFER CHOCOLATE",           "BENG BENG WAFER CHOCO 12 X 25 X 22GM - CHOCO VARIANT"),
    ("SKU020", "BENG BENG WAFER STRAWBERRY",          "BENG BENG WAFER STRAWBERRY 12 X 25 X 22GM - BERRY VARIANT"),
    ("SKU021", "BENG BENG PARTY PACK",                "BENG BENG PARTY PACK 6 X 10 X 25G - SHARING PACK"),
    ("SKU022", "BENG BENG CRISPY ROLL",               "BENG BENG CRISPY ROLL 12 X 20 X 18GM - ROLL FORMAT"),
    ("SKU023", "BENG BENG JUMBO WAFER",               "BENG BENG JUMBO 8 X 12 X 45GM - JUMBO SINGLE SERVE"),
    ("SKU024", "BENG BENG ASSORTED GIFT BOX",         "BENG BENG ASSORTED 3 FLAVOUR GIFT BOX 12 X 150GM - FESTIVE"),
    ("SKU025", "BENG BENG WAFER MINI BOX",            "BENG BENG MINI BOX 48 PCS X 11GM - MINI IMPULSE BOX"),
    ("SKU026", "KOPIKO CAPPUCCINO SACHET 10S",        "KOPIKO CAPPUCCINO 10 SACHET PACK X 3.5G - TRIAL PACK"),
    ("SKU027", "KOPIKO CAPPUCCINO JAR 100S",          "KOPIKO CAPPUCCINO JAR 100 PCS X 3.5G - BULK JAR"),
    ("SKU028", "KOPIKO BLANCA CANDY",                 "KOPIKO BLANCA 28 PCH X 120 X 3.5G - WHITE COFFEE VARIANT"),
    ("SKU029", "KOPIKO CAPPUCCINO DISPLAY TIN",       "KOPIKO CAPPUCCINO TIN 150 PCS X 3.5G - RETAIL DISPLAY TIN"),
    ("SKU030", "KOPIKO CAPPU EXTRA STRONG",           "KOPIKO CAPPU STRONG 12 X 230 X 3.5GR - EXTRA STRONG"),
    ("SKU031", "KOPIKO CAPPU DECAF",                  "KOPIKO CAPPU DECAF 12 X 230 X 3.5GR - DECAF VARIANT"),
    ("SKU032", "KOPIKO CAPPU + BENG BENG COMBO",      "KOPIKO CAPPU 4X650X3.5G + BENG BENG 20PCS COMBO - NEW COMBO"),
    ("SKU033", "KOPIKO CAPPUCCINO GIFT TIN",          "KOPIKO CAPPUCCINO GIFT TIN 200 PCS X 3.5G - PREMIUM GIFTING"),
    ("SKU034", "MALKIST SUGAR MINI PACK",             "MALKIST SUGAR MINI 24 PCS X 36GM - IMPULSE MINI"),
    ("SKU035", "MALKIST SUGAR + CREAM",               "MALKIST SUGAR CREAM 16 X 30 X 20G - CREAM FILLED"),
    ("SKU036", "MALKIST SUGAR FAMILY PACK",           "MALKIST SUGAR FAMILY 10 X 6 X 150GM - FAMILY VALUE"),
    ("SKU037", "MALKIST SUGAR THIN CRACKER",          "MALKIST SUGAR THIN 24 X 20 X 25GM - THIN LIGHT CRACKER"),
    ("SKU038", "MALKIST SUGAR MULTIPACK",             "MALKIST SUGAR MULTIPACK 5 X 10 X 80GM - SCHOOL MULTIPACK"),
    ("SKU039", "MALKIST SUGAR GIFT BOX",              "MALKIST SUGAR GIFT 3 X 200GM - PREMIUM FESTIVE GIFT"),
    ("SKU040", "MALKIST SUGAR BULK SHIPPER",          "MALKIST SUGAR SHIPPER 100 PCS X 80GM - BULK TRADE CASE"),
]


def generate_demand_excel(distributor_id: str, recommendations: list) -> str:
    """
    Generates an Excel file matching the DEMAND.xlsx format exactly:
    Columns: sku_id | sku_name | sku_description | Quantity
    All 40 SKUs included. Recommended SKUs are highlighted in yellow.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    # ── Header row ─────────────────────────────────────────────────────────────
    headers = ["sku_id", "sku_name", "sku_description", "Quantity"]
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")

    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    # ── Recommended SKU names (for highlighting) ────────────────────────────
    rec_names = [r.upper() for r in recommendations]
    highlight_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")

    # ── Data rows ───────────────────────────────────────────────────────────────
    for row_idx, (sku_id, sku_name, sku_desc) in enumerate(ALL_SKUS, start=2):
        ws.cell(row=row_idx, column=1, value=sku_id)
        ws.cell(row=row_idx, column=2, value=sku_name)
        ws.cell(row=row_idx, column=3, value=sku_desc)
        qty_cell = ws.cell(row=row_idx, column=4, value=None)

        # Highlight recommended SKUs in yellow
        if sku_name.upper() in rec_names:
            for col in range(1, 5):
                ws.cell(row=row_idx, column=col).fill = highlight_fill

    # ── Column widths ───────────────────────────────────────────────────────────
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 35
    ws.column_dimensions["C"].width = 55
    ws.column_dimensions["D"].width = 15

    # ── Save ────────────────────────────────────────────────────────────────────
    attachments_dir = os.getenv("ATTACHMENTS_DIR", "attachments")
    os.makedirs(attachments_dir, exist_ok=True)
    file_path = os.path.join(attachments_dir, "DEMAND.xlsx")
    wb.save(file_path)

    return file_path


def send_email(
    to_email: str,
    subject: str,
    body: str,
    attachment_path: str = None,
) -> dict:
    try:
        msg = MIMEMultipart()
        msg["From"] = settings.EMAIL_ADDRESS
        msg["To"] = to_email
        msg["Subject"] = subject
        msg.attach(MIMEText(body, "plain"))

        if attachment_path and os.path.exists(attachment_path):
            with open(attachment_path, "rb") as f:
                part = MIMEBase("application", "octet-stream")
                part.set_payload(f.read())
            encoders.encode_base64(part)
            filename = os.path.basename(attachment_path)
            part.add_header(
                "Content-Disposition",
                f"attachment; filename={filename}",
            )
            msg.attach(part)

        context = ssl.create_default_context()
        with smtplib.SMTP_SSL(
            settings.SMTP_SERVER, settings.SMTP_PORT, context=context
        ) as server:
            server.login(settings.EMAIL_ADDRESS, settings.EMAIL_APP_PASSWORD)
            server.send_message(msg)

        return {"status": "success", "message": f"Email sent to {to_email}"}
    except Exception as e:
        return {"status": "error", "message": str(e)}


def send_bulk_emails(distributors: list) -> dict:
    results = []
    rec_service = SKURecommendationService()

    for distributor in distributors:
        distributor_id = distributor.get("id")
        email = distributor.get("email")

        try:
            # Get recommendations from FalkorDB + recommended_products table
            rec_result = rec_service.execute(distributor_id)
            recommendations = rec_result.get("recommendations", [])

            # Build personalised email body
            email_payload = build_demand_email(
                distributor_id=distributor_id,
                distributor_email=email,
                recommendations=recommendations,
            )

            # Generate Excel with all 40 SKUs, recommended ones highlighted
            attachment_path = generate_demand_excel(distributor_id, recommendations)

            # Send email with attachment
            result = send_email(
                to_email=email,
                subject=email_payload["subject"],
                body=email_payload["body"],
                attachment_path=attachment_path,
            )

            results.append({
                "distributor_id": distributor_id,
                "email": email,
                "recommendations_count": len(recommendations),
                "attachment": os.path.basename(attachment_path),
                "status": result["status"],
            })

        except Exception as e:
            results.append({
                "distributor_id": distributor_id,
                "email": email,
                "status": "error",
                "message": str(e),
            })

    return {"status": "completed", "results": results}